from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "deliverables" / "SIS-System-Design-Presentation.pptx"
NAVY = RGBColor(11, 18, 32)
PANEL = RGBColor(18, 29, 48)
LINE = RGBColor(39, 57, 80)
TEXT = RGBColor(231, 237, 247)
MUTED = RGBColor(151, 168, 190)
TEAL = RGBColor(119, 212, 197)
AMBER = RGBColor(255, 190, 105)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def textbox(slide, x, y, w, h, text, size=18, color=TEXT, bold=False,
            font="Aptos", align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.02)
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def panel(slide, x, y, w, h, fill=PANEL, line=LINE, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def base(title, section, number):
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    panel(slide, 0, 0, 13.333, 0.12, TEAL, TEAL, False)
    textbox(slide, 0.55, 0.34, 11.9, 0.25, section.upper(), 10, TEAL, True)
    textbox(slide, 0.55, 0.68, 12.1, 0.62, title, 27, TEXT, True)
    panel(slide, 0.55, 7.08, 12.2, 0.012, LINE, LINE, False)
    textbox(slide, 0.58, 7.14, 10.5, 0.2, "SIS SCALABILITY LAB  •  SYSTEM DESIGN", 8, MUTED, True)
    textbox(slide, 12.0, 7.12, 0.65, 0.22, f"{number:02d}", 9, MUTED, True, align=PP_ALIGN.RIGHT)
    return slide


def bullet_slide(title, section, number, bullets, takeaway=None):
    slide = base(title, section, number)
    y = 1.62
    for i, (head, body) in enumerate(bullets):
        panel(slide, 0.7, y, 11.95, 0.72)
        textbox(slide, 0.95, y + 0.11, 2.45, 0.42, head, 15, TEAL, True)
        textbox(slide, 3.3, y + 0.09, 8.95, 0.5, body, 13, TEXT)
        y += 0.88
    if takeaway:
        panel(slide, 0.7, 6.35, 11.95, 0.48, RGBColor(22, 48, 55), TEAL)
        textbox(slide, 0.9, 6.42, 11.5, 0.27, takeaway, 12, TEAL, True)
    return slide

# Slide 1: title
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
panel(slide, 0, 0, 13.333, 0.14, TEAL, TEAL, False)
textbox(slide, 0.8, 1.05, 11.5, 0.35, "UCS3513  •  SYSTEM DESIGN LABORATORY", 13, TEAL, True)
textbox(slide, 0.8, 1.75, 11.6, 1.35, "Scalable Student\nInformation System", 34, TEXT, True)
textbox(slide, 0.82, 3.55, 10.5, 0.75,
        "Correctness-first architecture, concurrent enrollment safety, and measured scaling trade-offs.",
        19, MUTED)
panel(slide, 0.82, 5.0, 3.6, 0.78, PANEL, LINE)
textbox(slide, 1.05, 5.18, 3.1, 0.32, "23 automated tests passing", 15, TEAL, True)
panel(slide, 4.62, 5.0, 3.6, 0.78, PANEL, LINE)
textbox(slide, 4.85, 5.18, 3.1, 0.32, "PostgreSQL concurrency test", 14, TEXT, True)
panel(slide, 8.42, 5.0, 3.95, 0.78, PANEL, LINE)
textbox(slide, 8.65, 5.18, 3.5, 0.32, "8 required components demonstrated", 13, TEXT, True)
textbox(slide, 0.82, 6.7, 11.5, 0.28, "Prototype scope and limits are stated explicitly; no production-readiness claim.", 11, MUTED)

# Slide 2: architecture
slide = base("One authoritative database, isolated experiments", "Architecture", 2)
cards = [
    (0.7, "CLIENTS", "Browser dashboard\nLocust workloads"),
    (4.55, "API LAYER", "FastAPI modular monolith\nValidation + token bucket"),
    (8.4, "SOURCE OF TRUTH", "PostgreSQL\nAcademic records + transactions"),
]
for x, head, body in cards:
    panel(slide, x, 2.1, 3.45, 1.35)
    textbox(slide, x + 0.18, 2.32, 3.1, 0.3, head, 13, TEAL, True, align=PP_ALIGN.CENTER)
    textbox(slide, x + 0.18, 2.76, 3.1, 0.5, body, 13, TEXT, align=PP_ALIGN.CENTER)
for x in [4.12, 7.97]:
    textbox(slide, x, 2.58, 0.35, 0.4, "→", 25, AMBER, True, align=PP_ALIGN.CENTER)
panel(slide, 1.0, 4.25, 11.25, 1.15, RGBColor(22, 34, 53), LINE)
textbox(slide, 1.25, 4.48, 2.25, 0.3, "SYSTEM-DESIGN LAB", 13, TEAL, True)
textbox(slide, 3.45, 4.43, 8.4, 0.65,
        "Cache · hashing · shard simulation · ID generator · short links · crawler · autocomplete · metrics",
        14, TEXT)
textbox(slide, 1.0, 5.8, 11.1, 0.55,
        "The in-memory components demonstrate mechanisms; they are not claimed as durable distributed services.",
        13, AMBER, True)

# Slide 3: domain + invariants
bullet_slide("The data model reflects academic reality", "Domain model", 3, [
    ("Course", "Catalog identity: code, title and credits."),
    ("Offering", "A course in a particular term/section, with capacity and registration status."),
    ("Enrollment", "Student + offering, status and idempotency key; database constraints prevent duplicates."),
    ("Academic records", "Attendance per session, bounded assessment grades, timetable events and notifications."),
], "Invariant: 0 ≤ enrolled_count ≤ capacity. The database—not the UI—decides the final seat.")

# Slide 4: concurrency
slide = base("The last seat is a transaction, not a UI check", "Correctness", 4)
panel(slide, 0.7, 1.75, 5.55, 2.45)
textbox(slide, 1.0, 1.98, 4.9, 0.35, "NAÏVE COUNT → INSERT", 15, AMBER, True)
textbox(slide, 1.0, 2.5, 4.8, 1.2,
        "Two requests can both read one remaining seat, then both insert. A transaction alone does not automatically serialize that check.",
        16, TEXT)
panel(slide, 6.65, 1.75, 5.95, 2.45, RGBColor(20, 43, 47), TEAL)
textbox(slide, 6.95, 1.98, 5.2, 0.35, "CONDITIONAL ATOMIC UPDATE", 15, TEAL, True)
textbox(slide, 6.95, 2.5, 5.15, 1.2,
        "Update only when status is OPEN and enrolled_count < capacity; insert in the same transaction; unique constraints remain the final guard.",
        16, TEXT)
panel(slide, 0.7, 4.65, 11.9, 1.15, PANEL, LINE)
textbox(slide, 0.95, 4.87, 2.3, 0.35, "LIVE TEST RESULT", 13, TEAL, True)
textbox(slide, 3.15, 4.82, 8.95, 0.45,
        "20 simultaneous students · 5 seats · 5 accepted · 15 rejected · 0 over-capacity",
        17, TEXT, True)
textbox(slide, 0.95, 6.1, 11.2, 0.35, "A retry uses an idempotency key; a failed insert rolls back the seat reservation.", 13, MUTED)

# Slide 5: workflows
bullet_slide("Academic workflows include validation", "Prototype", 5, [
    ("Identity", "Registration, PBKDF2 password hashing, login and short-lived signed bearer token."),
    ("Registration", "Discover offerings, enroll, retry idempotently and drop once to release a seat."),
    ("Attendance", "Requires active enrollment; duplicate student/offering/session rows are rejected."),
    ("Grades", "Marks cannot exceed assessment maximum; only enrolled students can be graded."),
    ("Timetable + alerts", "Room/time overlap validation and per-student notification read state."),
], "Role-based checks protect staff mutations and student records; token revocation/refresh and production identity controls remain future hardening.")

# Slide 6: required components
slide = base("Eight required components, each with a measurable job", "System-design components", 6)
items = [
    ("Rate limiter", "429 rate, accepted/rejected requests"),
    ("Sharding", "Per-shard records and readability"),
    ("Consistent hashing", "Moved-key fraction after node change"),
    ("Key-value cache", "Hit/miss and response latency"),
    ("Unique IDs", "IDs/second and collision checks"),
    ("URL shortener", "Creation and redirect latency"),
    ("Web crawler", "Pages crawled, errors, frontier size"),
    ("Autocomplete", "Prefix latency and result relevance"),
]
for i, (head, body) in enumerate(items):
    col = i % 2
    row = i // 2
    x = 0.72 + col * 6.08
    y = 1.62 + row * 1.18
    panel(slide, x, y, 5.8, 0.92)
    textbox(slide, x + 0.18, y + 0.12, 1.95, 0.28, head, 13, TEAL, True)
    textbox(slide, x + 2.12, y + 0.12, 3.4, 0.52, body, 12, TEXT)
textbox(slide, 0.85, 6.48, 11.5, 0.35,
        "Adding a fourth node moved 277/1,000 sampled keys (27.7%); the ring routes but does not migrate data.",
        12, AMBER, True)

# Slide 7: benchmark results
slide = base("Short baseline runs establish a starting point", "Measured results", 7)
metrics = [
    ("Course discovery", "35.82 req/s", "p50 5 ms  •  p95 10 ms"),
    ("Profile reads", "35.27 req/s", "p50 5 ms  •  p95 12 ms"),
    ("Account registration", "19.47 req/s", "p50 60 ms  •  p95 83 ms"),
    ("Autocomplete", "35.55 req/s", "p50 2 ms  •  p95 5 ms"),
]
for i, (head, rate, lat) in enumerate(metrics):
    x = 0.72 + (i % 2) * 6.05
    y = 1.72 + (i // 2) * 1.62
    panel(slide, x, y, 5.75, 1.28)
    textbox(slide, x + 0.22, y + 0.15, 2.3, 0.28, head, 13, MUTED, True)
    textbox(slide, x + 0.22, y + 0.48, 2.45, 0.4, rate, 21, TEAL, True)
    textbox(slide, x + 2.75, y + 0.53, 2.65, 0.34, lat, 12, TEXT)
panel(slide, 0.72, 5.25, 11.8, 0.9, RGBColor(45, 34, 25), AMBER)
textbox(slide, 0.98, 5.47, 11.25, 0.42,
        "10 users · 20 seconds · synthetic 1k students / 100 courses / 200 offerings · limiter disabled · 0 failures",
        12, AMBER, True)
textbox(slide, 0.85, 6.32, 11.4, 0.45,
        "Autocomplete cache OFF vs ON: p50/p95 5/10 ms vs 2/5 ms (one 20s run each); repeat before capacity claims.",
        10, MUTED)

# Slide 8: limitations
bullet_slide("What this prototype does not claim", "Trade-offs and failure modes", 8, [
    ("Auth + access", "PBKDF2/HMAC tokens; staff mutations and student record reads enforce role/self-access checks."),
    ("Local state", "Rate limiter, cache, shard store and metrics reset when the process restarts; short-link mappings persist in PostgreSQL."),
    ("Shard topology", "Changing the ring does not move, replicate or recover stored records."),
    ("Cache", "Autocomplete cache is TTL-based and can be briefly stale; never use it for seat/grade decisions."),
    ("Crawler", "Same-host public HTTPS only; live DNS fetch unavailable here, so network behavior needs validation."),
], "Use synthetic records only. Do not expose this coursework prototype to the public internet.")

# Slide 9: conclusion
slide = base("A defensible design is one you can explain and test", "Conclusion", 9)
panel(slide, 0.75, 1.78, 5.7, 2.25, RGBColor(20, 43, 47), TEAL)
textbox(slide, 1.05, 2.03, 5.1, 0.4, "WHAT WE HAVE PROVED", 15, TEAL, True)
textbox(slide, 1.05, 2.58, 4.95, 1.1,
        "Database-enforced seat reservation survives a concurrent capacity test. Academic workflows and component behaviors have automated coverage.",
        17, TEXT)
panel(slide, 6.75, 1.78, 5.75, 2.25)
textbox(slide, 7.05, 2.03, 5.1, 0.4, "NEXT ENGINEERING STEPS", 15, AMBER, True)
textbox(slide, 7.05, 2.58, 4.95, 1.1,
        "Add token revocation, repeat longer benchmarks with resource sampling, use shared Redis state, and implement real shard migration before scaling claims.",
        16, TEXT)
textbox(slide, 0.95, 4.7, 11.3, 0.7,
        "Differentiator: an Academic Safety Case—showing the transactional constraints behind registration outcomes, not merely adding features.",
        19, TEXT, True)
textbox(slide, 0.95, 6.0, 11.3, 0.38, "Questions", 22, TEAL, True)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"Created {OUT} ({len(prs.slides)} slides)")
